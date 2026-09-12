"""Live-network smoke test for the SBI adapter (spec §9: manually triggered, skips if network
is unavailable). Proves the multi-AMC pattern generalizes beyond PPFAS: a JS-SPA disclosure
page whose real download links live behind a POST endpoint (found via one-time Playwright
network capture, reproduced here with plain httpx — no browser at runtime), and a combined
one-workbook-per-month / one-sheet-per-scheme file layout, distinct from PPFAS's one-file-
per-scheme layout.
"""
from datetime import date

import httpx
import openpyxl
import pytest

from indian_mf_mcp.ingest.amc_adapters.sbi import SBIAdapter
from indian_mf_mcp.ingest.portfolio_ingest import ingest_scheme_portfolios
from indian_mf_mcp.store import repository as repo
from indian_mf_mcp.store.db import get_connection
from indian_mf_mcp.tools.get_fund_portfolio import get_fund_portfolio

SCHEME_ID = "scheme-sbi-flexicap-test"
AMC_ID = "amc-sbi"
SCHEME_HINT = "SBI Flexicap Fund"


def _network_available() -> bool:
    try:
        httpx.get("https://www.sbimf.com/", timeout=10)
        return True
    except Exception:
        return False


def _sheet_resolver(raw: bytes, scheme_hint: str) -> str | None:
    wb = openpyxl.load_workbook(__import__("io").BytesIO(raw), read_only=True, data_only=True)
    if "Index" not in wb.sheetnames:
        return None
    index_rows = list(wb["Index"].iter_rows(values_only=True))
    return SBIAdapter.find_sheet_code(index_rows, scheme_hint)


@pytest.mark.skipif(not _network_available(), reason="network unavailable")
def test_sbi_end_to_end_live(tmp_path):
    conn = get_connection(tmp_path / "sbi_live.db")
    repo.upsert_amc(conn, AMC_ID, "SBI Funds Management Limited")
    repo.upsert_scheme(conn, SCHEME_ID, AMC_ID, SCHEME_HINT, "Open Ended Schemes",
                        "Equity Scheme", "Flexi Cap Fund", date.today().isoformat())
    conn.commit()

    adapter = SBIAdapter()
    stats = ingest_scheme_portfolios(
        conn, adapter, SCHEME_ID, SCHEME_HINT, since=date(2026, 6, 1),
        sheet_resolver=_sheet_resolver,
    )
    conn.commit()

    assert stats["documents_found"] > 0
    assert stats["ingested"] > 0
    assert stats["skipped_sheet_not_found"] == 0
    assert stats["skipped_reconciliation_failed"] == 0

    out = get_fund_portfolio(conn, [SCHEME_ID], sections=["holdings", "concentration"])
    data = out[SCHEME_ID]["data"]
    assert len(data["holdings"]["v"]) > 30
    icici = next((h for h in data["holdings"]["v"] if "ICICI Bank" in h["instrument_name"]), None)
    assert icici is not None
    assert icici["isin"] == "INE090A01021"
    assert 0 < data["concentration"]["v"]["top10_pct"] < 1  # fractional, normalised from % points

    conn.close()
