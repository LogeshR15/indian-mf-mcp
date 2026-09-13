"""Unit tests for ingest/amfi_caplist.py — AMFI cap-list parse, store, and lookup."""
from __future__ import annotations

import io
import sqlite3
from datetime import date
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from indian_mf_mcp.ingest.amfi_caplist import (
    _classify_cap_raw,
    _derive_effective_date,
    _parse_caplist_xlsx,
    compute_market_cap_allocation,
    get_market_cap,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_xlsx_section_layout(entries: list[tuple[str, str, str]]) -> bytes:
    """Build a minimal cap-list XLSX using the section-header layout.

    entries = [(cap_label, isin, company_name), ...]
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    current_section = None
    row = 1
    for cap_label, isin, name in entries:
        if cap_label != current_section:
            ws.cell(row=row, column=1, value=f"{cap_label.title()} Cap Companies")
            current_section = cap_label
            row += 1
        ws.cell(row=row, column=1, value=isin)
        ws.cell(row=row, column=2, value=name)
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_xlsx_flat_layout(entries: list[tuple[str, str, str]]) -> bytes:
    """Build a minimal cap-list XLSX using the flat table layout.

    entries = [(isin, company_name, market_cap_label), ...]
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="ISIN")
    ws.cell(row=1, column=2, value="Company Name")
    ws.cell(row=1, column=3, value="Market Cap")
    for r, (isin, name, cap) in enumerate(entries, start=2):
        ws.cell(row=r, column=1, value=isin)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=cap)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _in_memory_db():
    """Create an in-memory SQLite DB with the full schema."""
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    with open("src/indian_mf_mcp/store/schema.sql") as f:
        conn.executescript(f.read())
    return conn


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestClassifyCapRaw:
    def test_large(self):
        assert _classify_cap_raw("Large Cap Companies") == "large"

    def test_mid(self):
        assert _classify_cap_raw("Mid-Cap") == "mid"

    def test_small(self):
        assert _classify_cap_raw("Small cap") == "small"

    def test_unknown_returns_none(self):
        assert _classify_cap_raw("Debt Fund") is None

    def test_case_insensitive(self):
        assert _classify_cap_raw("LARGE CAP") == "large"


class TestDeriveEffectiveDate:
    def test_returns_jan_or_jul(self):
        eff = _derive_effective_date()
        assert eff.endswith("-01-01") or eff.endswith("-07-01")

    def test_format_is_iso(self):
        eff = _derive_effective_date()
        date.fromisoformat(eff)  # should not raise


class TestParseCaplistXlsx:
    def test_section_layout_large(self):
        raw = _make_xlsx_section_layout([
            ("large", "INE040A01034", "HDFC Bank"),
            ("large", "INE009A01021", "Infosys"),
            ("mid", "INE030A01027", "Mid Corp"),
        ])
        eff_date, entries = _parse_caplist_xlsx(raw)
        isins = {isin: cap for isin, _, cap in entries}
        assert isins["INE040A01034"] == "large"
        assert isins["INE009A01021"] == "large"
        assert isins["INE030A01027"] == "mid"

    def test_flat_layout(self):
        # Use real-format ISINs (12 chars: IN + 10 alphanumeric) so _ISIN_RE matches
        raw = _make_xlsx_flat_layout([
            ("INE001A01023", "Alpha Corp", "Large Cap"),
            ("INE002B01023", "Beta Corp", "Mid Cap"),
            ("INE003C01023", "Gamma Corp", "Small Cap"),
        ])
        eff_date, entries = _parse_caplist_xlsx(raw)
        isins = {isin: cap for isin, _, cap in entries}
        assert isins["INE001A01023"] == "large"
        assert isins["INE002B01023"] == "mid"
        assert isins["INE003C01023"] == "small"

    def test_effective_date_fallback(self):
        raw = _make_xlsx_section_layout([("large", "INE001A01023", "Test Co")])
        eff_date, _ = _parse_caplist_xlsx(raw)
        # No date in the file → derived from current date
        date.fromisoformat(eff_date)  # should be valid ISO

    def test_empty_file_returns_empty_entries(self):
        wb = openpyxl.Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        eff_date, entries = _parse_caplist_xlsx(buf.getvalue())
        assert entries == []


class TestGetMarketCap:
    def test_lookup_found(self):
        conn = _in_memory_db()
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE040A01034', 'large', '2026-07-01', NULL)"
        )
        result = get_market_cap(conn, "INE040A01034", "2026-08-31")
        assert result == "large"

    def test_lookup_not_found(self):
        conn = _in_memory_db()
        result = get_market_cap(conn, "INE999Z99999", "2026-08-31")
        assert result is None

    def test_point_in_time_uses_correct_version(self):
        """Should use the cap-list whose effective_date is <= the query date."""
        conn = _in_memory_db()
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE001A01023', 'mid', '2026-01-01', NULL)"
        )
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE001A01023', 'large', '2026-07-01', NULL)"
        )
        # Query before July revision — should return 'mid'
        assert get_market_cap(conn, "INE001A01023", "2026-06-30") == "mid"
        # Query after July revision — should return 'large'
        assert get_market_cap(conn, "INE001A01023", "2026-08-01") == "large"

    def test_future_cap_list_not_used(self):
        """A cap list from 2027 should NOT be used for a 2026 portfolio."""
        conn = _in_memory_db()
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE001A01023', 'small', '2027-01-01', NULL)"
        )
        # Querying for 2026 → no matching cap list → None
        assert get_market_cap(conn, "INE001A01023", "2026-12-31") is None


class TestComputeMarketCapAllocation:
    def _make_holding(self, isin, pct, asset_class="equity"):
        return {
            "isin": isin, "pct_nav": pct, "asset_class": asset_class,
            "instrument_name": "Test",
        }

    def test_with_cap_list(self):
        conn = _in_memory_db()
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE001A01023', 'large', '2026-07-01', NULL)"
        )
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE002B01023', 'mid', '2026-07-01', NULL)"
        )
        holdings = [
            self._make_holding("INE001A01023", 40.0),
            self._make_holding("INE002B01023", 30.0),
        ]
        result = compute_market_cap_allocation(conn, holdings, "2026-08-31")
        assert result["large_pct"] == pytest.approx(40.0)
        assert result["mid_pct"] == pytest.approx(30.0)
        assert result["small_pct"] == pytest.approx(0.0)
        assert result["cap_list_date"] == "2026-07-01"
        assert result["n_classified"] == 2

    def test_no_cap_list_returns_none_values(self):
        conn = _in_memory_db()
        holdings = [self._make_holding("INE001A01023", 40.0)]
        result = compute_market_cap_allocation(conn, holdings, "2026-08-31")
        assert result["large_pct"] is None
        assert "update-caplist" in result["caveat"]

    def test_excludes_cash_and_debt(self):
        conn = _in_memory_db()
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE001A01023', 'large', '2026-07-01', NULL)"
        )
        holdings = [
            self._make_holding("INE001A01023", 40.0, "equity"),
            self._make_holding("INE002B01023", 20.0, "debt"),    # excluded
            self._make_holding(None, 5.0, "cash"),               # excluded
        ]
        result = compute_market_cap_allocation(conn, holdings, "2026-08-31")
        assert result["large_pct"] == pytest.approx(40.0)
        assert result["n_classified"] == 1

    def test_unclassified_isin(self):
        conn = _in_memory_db()
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE001A01023', 'large', '2026-07-01', NULL)"
        )
        holdings = [
            self._make_holding("INE001A01023", 40.0),
            self._make_holding("INE_UNKNOWN", 20.0),  # not in cap list
        ]
        result = compute_market_cap_allocation(conn, holdings, "2026-08-31")
        assert result["large_pct"] == pytest.approx(40.0)
        assert result["unclassified_pct"] == pytest.approx(20.0)

    def test_null_isin_counted_as_unclassified(self):
        conn = _in_memory_db()
        conn.execute(
            "INSERT INTO isin_market_cap (isin, market_cap, effective_date, caplist_doc_id) "
            "VALUES ('INE001A01023', 'large', '2026-07-01', NULL)"
        )
        holdings = [
            self._make_holding(None, 10.0, "equity"),   # no ISIN → unclassified
            self._make_holding("INE001A01023", 40.0),
        ]
        result = compute_market_cap_allocation(conn, holdings, "2026-08-31")
        assert result["unclassified_pct"] == pytest.approx(10.0)
        assert result["large_pct"] == pytest.approx(40.0)
