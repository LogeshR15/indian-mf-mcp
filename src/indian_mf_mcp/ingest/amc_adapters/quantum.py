"""Quantum Mutual Fund adapter (quantumamc.com — do not confuse with "quant Mutual Fund",
quantmutual.com, which is a different AMC with its own adapter at `quant.py`).

Ground-truthed live 2026-09-13. AMFI's own registry names the disclosure page as
``https://www.quantumamc.com/portfolio/combined/-1/1/0/0`` — this is not a SPA and needs no
API reverse-engineering at all: a plain `httpx.get()` with this project's honest User-Agent
returns a clean `200` with the *entire* requested listing already rendered server-side as a
static `<ul>` of download links, each carrying its own
``onclick="GTMcodeforxml('<url>', '<page-url>', '<title>', '<subtitle>')"`` analytics call —
that literal is the cheapest possible source of (file URL, human title) pairs, scraped with a
single regex; no inline JS array, no `__NEXT_DATA__`, no AJAX endpoint, no Playwright capture
was needed anywhere in this adapter.

**URL shape**: the four numeric path segments in the registered URL,
``/portfolio/combined/{scheme_id}/{portfolio_type}/{year}/{month}``, are exactly what they
look like — filter parameters for a server-side query, not routing cruft:

  - `{scheme_id}`: `-1` means "All Funds". Verified live that changing this to a specific
    scheme's internal id (e.g. `16` for "Quantum Flexi Cap Fund", read off the page's own
    `#ddlProductScheme` `<select>`) returns **the same list of combined workbooks** — this
    AMC never publishes a per-scheme file, only one combined workbook per period covering
    every scheme, so the scheme filter is cosmetic for our purposes and this adapter always
    requests `-1`.
  - `{portfolio_type}`: `1` = Monthly, `2` = Weekly, `3` = Fortnightly (from the page's own
    radio-button group, `name="PortfolioType"`). Only Monthly is used here.
  - `{year}` / `{month}`: `0/0` (both zero) is what the page's own default view uses and
    returns the most recent ~20 months spanning a year boundary (verified: a fetch on
    2026-09-13 returned Jan 2025 through Aug 2026, 20 entries) — convenient for humans
    clicking through the site, but an unreliable way to enumerate a specific range
    programmatically. **This is the history lever**: passing a specific `{year}` with
    `month=0` returns every month *actually published* for that calendar year in one request
    (verified: `year=2022` → 12 entries Jan-Dec 2022; `year=2015` → 6 entries, Jul-Dec 2015
    only, i.e. before that the AMC evidently didn't publish under this URL scheme; `year=2011`
    and `year=2012` → 0 entries). So history is reached by iterating `since.year` through the
    current year, one GET per year (at most ~15 requests for the full archive), never by
    guessing per-month URLs.

**Where the files live**: every entry's actual document is `https://www.quantumamc.com/FileCDN
/FactSheet/<uuid>.xlsx` — same host as the listing page (no separate CDN/S3 host, unlike HDFC
or Union), served with the honest User-Agent as a clean `200
application/vnd.openxmlformats-officedocument.spreadsheetml.sheet`. The filename itself is an
opaque UUID and carries no date or scheme information — the only place the as-of date is
recoverable is the anchor's own analytics-call title text (e.g. `"August 2026 - All Funds"`),
which is what this adapter parses for `as_of_date`, exactly as the existing `quant.py` and
`union.py` adapters do for their own title/filename-unreliable AMCs.

**File shape**: one **combined workbook** per month covering every Quantum scheme (SBI/Tata/
Sundaram-style, not PPFAS/Union/HDFC one-file-per-scheme). An `Index` sheet maps each scheme
to its own per-scheme sheet, and each scheme sheet is named after a short internal acronym
("QLTEVF" for Quantum Value Fund, "QLF" for Quantum Liquid Fund, "QSCAPF" for Quantum Small
Cap Fund, etc. — 15 sheets total as of August 2026, one per live scheme).

**Quirk requiring a *local* (not shared) sheet-resolver**: the shared
`combined_workbook.find_sheet_code` helper looks for a header cell containing the substring
"scheme name" (or "fund name") to find the name column — but Quantum's own Index sheet header
row reads **"Scheme Full Name" / "Scheme Code"**, and `"scheme full name"` does not contain
`"scheme name"` as a contiguous substring (there's "full" in the middle), so the shared helper
silently returns `None` for every scheme here. Rather than loosen the shared keyword matching
project-wide (risking a false match on some other AMC's wording), `_resolve_sheet` below is a
small local variant that additionally recognises "scheme full name" — kept local to this file
per CONTRIBUTING.md's preference for AMC-specific repairs over shared-parser changes. (A
shared-parser diff, if the maintainer wants it instead: in `combined_workbook.find_sheet_code`,
change the name-column test from `"fund name" in t or "scheme name" in t` to also match
`"scheme full name" in t` — a strict superstring addition that cannot regress any AMC whose
header already matches "scheme name"/"fund name".)

**Parser fit**: each scheme sheet uses the standard "Sr.No. / Name of Instrument / ISIN /
Industry+ / Quantity / Market/Fair Value / % to NAV" header row, values already fractional
(e.g. `0.0543` = 5.43%, not percentage-points), and a plain "Grand Total" row summing to
`1.0`. Ground-truthed on Quantum Value Fund, August 2026: exact 100% reconciliation, 38
holdings, zero warnings — parses against `xlsx_portfolio.py` with **zero shared-parser
changes**.

**Coverage / quirks**:
  - Only `DocType.MONTHLY_PORTFOLIO` is implemented (portfolio_type=1); weekly/fortnightly
    (2/3) use the identical URL shape and could be added the same way if needed.
  - The scheme lineup itself has changed over the archive's life (schemes merge, rename, and
    launch — e.g. "Quantum Small Cap Fund" and "Quantum Multi Asset Allocation Fund" are
    recent additions), so `_resolve_sheet`'s fuzzy name matching (exact match preferred,
    substring match as fallback) is required, not merely convenient, the same way it is for
    every other combined-workbook AMC here.
  - No CAPTCHA, no WAF, no anti-bot challenge encountered anywhere in this adapter — this is
    a clean, fully public, un-gated disclosure flow end to end.
"""
from __future__ import annotations

import calendar
import io
import re
from datetime import date, datetime

import httpx
import openpyxl

from indian_mf_mcp import config
from indian_mf_mcp.ingest.amc_adapters.base import DocType, DocumentRef

BASE_URL = "https://www.quantumamc.com"
LISTING_PATH = "/portfolio/combined/-1/{portfolio_type}/{year}/0"
_MONTHLY_PORTFOLIO_TYPE = 1

_ENTRY_RE = re.compile(
    r"GTMcodeforxml\('([^']+)',\s*'[^']*',\s*'([^']+)',\s*'([^']+)'\)"
)
_TITLE_DATE_RE = re.compile(r"([A-Za-z]+)\s+(\d{4})")


def _years_since(since: date) -> list[int]:
    return list(range(since.year, date.today().year + 1))


def _resolve_sheet(raw: bytes, scheme_hint: str) -> str | None:
    """Local variant of `combined_workbook.find_sheet_code`: Quantum's own Index sheet header
    reads "Scheme Full Name" / "Scheme Code", which the shared helper's "scheme name"
    substring test does not match (see module docstring). Same fuzzy-match semantics as the
    shared helper otherwise: exact (case-insensitive) name match preferred, substring
    containment as a fallback for schemes not spelled identically to `scheme_hint`."""
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if "Index" not in wb.sheetnames:
        return None
    rows = list(wb["Index"].iter_rows(values_only=True))

    name_col = code_col = header_idx = None
    for i, row in enumerate(rows):
        texts = [str(c).strip().lower() if isinstance(c, str) else "" for c in row]
        nc = next((idx for idx, t in enumerate(texts)
                    if "scheme full name" in t or "fund name" in t or "scheme name" in t), None)
        cc = next((idx for idx, t in enumerate(texts) if "scheme code" in t or "fund code" in t), None)
        if nc is not None and cc is not None:
            name_col, code_col, header_idx = nc, cc, i
            break
    if header_idx is None:
        return None

    target = scheme_hint.strip().lower()
    best = None
    for row in rows[header_idx + 1:]:
        if len(row) <= max(name_col, code_col):
            continue
        name, code = row[name_col], row[code_col]
        if not isinstance(name, str) or not code:
            continue
        name_l = name.strip().lower()
        if name_l == target:
            return str(code)
        if (target in name_l or name_l in target) and best is None:
            best = str(code)
    return best


class QuantumAdapter:
    amc_id = "amc-quantum"

    def list_documents(self, doc_type: DocType, since: date,
                        scheme_hint: str | None = None,
                        client: httpx.Client | None = None) -> list[DocumentRef]:
        if doc_type != DocType.MONTHLY_PORTFOLIO:
            return []
        headers = {"User-Agent": config.USER_AGENT}
        get = client.get if client is not None else httpx.get

        refs: list[DocumentRef] = []
        seen = set()
        for year in _years_since(since):
            path = LISTING_PATH.format(portfolio_type=_MONTHLY_PORTFOLIO_TYPE, year=year)
            resp = get(BASE_URL + path, headers=headers, timeout=45, follow_redirects=True)
            resp.raise_for_status()
            html = resp.text

            for url, title, _subtitle in _ENTRY_RE.findall(html):
                if not url.lower().split("?")[0].endswith(".xlsx"):
                    continue
                m = _TITLE_DATE_RE.search(title.strip())
                if not m:
                    continue
                month_name, year_str = m.groups()
                try:
                    month_start = datetime.strptime(f"{month_name} {year_str}", "%B %Y").date()
                except ValueError:
                    continue
                last_day = calendar.monthrange(month_start.year, month_start.month)[1]
                as_of = month_start.replace(day=last_day)
                if as_of < since:
                    continue
                key = (url.split("?")[0], as_of)
                if key in seen:
                    continue
                seen.add(key)
                refs.append(DocumentRef(url=url, doc_type=DocType.MONTHLY_PORTFOLIO,
                                         as_of_date=as_of, scheme_hint=scheme_hint))
        refs.sort(key=lambda r: r.as_of_date)
        return refs

    def fetch(self, ref: DocumentRef, client: httpx.Client | None = None) -> bytes:
        headers = {"User-Agent": config.USER_AGENT}
        get = client.get if client is not None else httpx.get
        resp = get(ref.url, headers=headers, follow_redirects=True, timeout=90)
        resp.raise_for_status()
        return resp.content
