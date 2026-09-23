"""Shared helper for AMCs that publish one combined workbook per month covering every scheme,
with an "Index"-style sheet mapping scheme names to per-scheme sheet codes (SBI, Motilal
Oswal, Tata, Nippon India so far). Column positions for the index sheet vary by AMC (SBI:
code then name; Motilal Oswal/Tata: name then code) — detected from the index sheet's own
header row rather than assumed, same principle as the main holdings parser. Nippon India's
Index sheet has no header row at all (just (code, name) pairs from row 0) — handled by a
headerless fallback that recognises the shape directly.

Format-agnostic: some combined-workbook AMCs' archives include genuine legacy .xls (BIFF)
months alongside newer .xlsx ones (e.g. 360 ONE's 2018-2020 files) — sheet lookup is done via
_open_workbook() below, which dispatches on sniff() the same way xlsx_portfolio.py's
parse_portfolio_xlsx/parse_portfolio_xls do, so a caller doesn't need to know the format
before asking "which sheet is this scheme in?".
"""
from __future__ import annotations

import io
from typing import Callable

import openpyxl
import xlrd

from indian_mf_mcp.ingest.amc_adapters.base import normalize_scheme_name
from indian_mf_mcp.parsers.sniff import FormatKind, sniff
from indian_mf_mcp.parsers.xlsx_portfolio import (
    parse_portfolio_xls, parse_portfolio_xlsx, xls_sheet_to_rows,
)

# A sheet code is short and has no spaces (e.g. "ME", "YO08", "144D"); a scheme name is a
# much longer free-text string. Used only when no header row can be found at all.
_MAX_CODE_LEN = 10


def _open_workbook(raw: bytes) -> tuple[list[str], Callable[..., list[tuple]]] | None:
    """Returns (sheet_names, rows_for), where rows_for(name, limit=None) reads a sheet's rows
    (optionally just the first `limit` of them — find_sheet_by_title only ever needs row 1,
    and re-reading every row of every sheet in a large combined workbook just to check its
    title would be a real performance regression, so the row-1-only optimisation the openpyxl
    read_only iterator gave for free is preserved explicitly here). Returns None if the format
    isn't one either parse_portfolio_xlsx or parse_portfolio_xls would accept — callers should
    already only reach here after a format check, matching portfolio_ingest.py's own
    sniff-then-dispatch gate, but staying honest (None, not an exception) if that changes."""
    fmt = sniff(raw)
    if fmt == FormatKind.XLSX:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)

        def rows_for_xlsx(name: str, limit: int | None = None) -> list[tuple]:
            ws = wb[name]
            if limit is not None:
                return list(ws.iter_rows(max_row=limit, values_only=True))
            return list(ws.iter_rows(values_only=True))

        return wb.sheetnames, rows_for_xlsx
    if fmt == FormatKind.XLS_BIFF:
        wb = xlrd.open_workbook(file_contents=raw)
        names = wb.sheet_names()

        def rows_for_xls(name: str, limit: int | None = None) -> list[tuple]:
            all_rows = xls_sheet_to_rows(wb.sheet_by_name(name))
            return all_rows[:limit] if limit is not None else all_rows

        return names, rows_for_xls
    return None


def _match(rows, header_idx, name_col, code_col, scheme_hint) -> str | None:
    target = normalize_scheme_name(scheme_hint)
    best = None
    for row in rows[header_idx + 1 :]:
        if len(row) <= max(name_col, code_col):
            continue
        name, code = row[name_col], row[code_col]
        if not isinstance(name, str) or not code:
            continue
        name_l = normalize_scheme_name(name)
        if not name_l:
            continue
        if name_l == target:
            return str(code)
        if (target in name_l or name_l in target) and best is None:
            best = str(code)
    return best


def find_sheet_code(raw: bytes, scheme_hint: str, index_sheet_name: str = "Index") -> str | None:
    opened = _open_workbook(raw)
    if opened is None:
        return None
    sheet_names, rows_for = opened
    if index_sheet_name not in sheet_names:
        return None
    rows = rows_for(index_sheet_name)

    name_col = code_col = None
    header_idx = None
    for i, row in enumerate(rows):
        texts = [str(c).strip().lower() if isinstance(c, str) else "" for c in row]
        nc = next((idx for idx, t in enumerate(texts) if "fund name" in t or "scheme name" in t), None)
        cc = next((idx for idx, t in enumerate(texts) if "fund code" in t or "scheme code" in t
                    or "short code" in t or "short name" in t or "acronym" in t), None)
        if nc is not None and cc is not None:
            name_col, code_col, header_idx = nc, cc, i
            break

    if header_idx is not None:
        return _match(rows, header_idx, name_col, code_col, scheme_hint)

    # Headerless fallback (e.g. Nippon India): find the first row that looks like a
    # (short code, long name) pair and treat everything from there as data, no header to skip.
    for i, row in enumerate(rows):
        if len(row) < 2:
            continue
        code, name = row[0], row[1]
        if (isinstance(code, str) and isinstance(name, str) and " " not in code.strip()
                and 0 < len(code.strip()) <= _MAX_CODE_LEN and len(name.strip()) > len(code.strip())):
            return _match(rows, i - 1, 1, 0, scheme_hint)  # header_idx=i-1 so row i itself is included
    return None


def find_sheet_by_title(raw: bytes, scheme_hint: str, exclude_sheet_names: tuple[str, ...] = ()) -> str | None:
    """For AMCs with no separate Index/lookup sheet at all (e.g. Franklin Templeton): the
    sheet name itself IS the per-scheme code, and each sheet's own row 1 carries the full
    scheme name in its own first populated cell. Scans every sheet's title row instead of a
    shared lookup table."""
    opened = _open_workbook(raw)
    if opened is None:
        return None
    sheet_names, rows_for = opened
    target = normalize_scheme_name(scheme_hint)
    best = None
    for sheet_name in sheet_names:
        if sheet_name in exclude_sheet_names:
            continue
        rows = rows_for(sheet_name, limit=1)
        if not rows:
            continue
        row0 = rows[0]
        title = next((c for c in row0 if isinstance(c, str) and c.strip()), None)
        if title is None:
            continue
        title_l = normalize_scheme_name(title)
        if not title_l:
            continue
        if title_l == target:
            return sheet_name
        if (target in title_l or title_l in target) and best is None:
            best = sheet_name
    return best


# How far down a sheet to look for the scheme's name in its title block.
_TITLE_SCAN_ROWS = 8


def find_sole_holdings_sheet(raw: bytes, scheme_hint: str) -> str | None:
    """For ONE-FILE-PER-SCHEME AMCs whose file nonetheless has extra sheets: HDFC adds a
    "Derivative<code>" disclosure sheet, Taurus a "<code> Performance" table, HSBC "Notes" and
    "Disclaimer", Invesco a top-10 summary, DSP an empty "Sheet1", ICICI a "Derivative" sheet.
    Returns the holdings sheet only when that is provable rather than a guess:

      1. exactly ONE sheet parses to holdings with at least one ISIN — every other sheet
         yields none; and
      2. that sheet's own title block names this scheme (normalised match).

    Anything else — two candidate sheets (a genuine combined workbook), zero, or a title that
    doesn't name the scheme — returns None, and the caller keeps treating the file as
    ambiguous. Parsing every sheet is deliberate: "the first sheet" or "the biggest sheet"
    would be exactly the guess portfolio_ingest's ambiguity gate exists to prevent."""
    fmt = sniff(raw)
    parse_fn = {FormatKind.XLSX: parse_portfolio_xlsx,
                FormatKind.XLS_BIFF: parse_portfolio_xls}.get(fmt)
    opened = _open_workbook(raw)
    if parse_fn is None or opened is None:
        return None
    sheet_names, rows_for = opened

    candidates = [name for name in sheet_names
                  if any(h.isin for h in parse_fn(raw, sheet_name=name).holdings)]
    if len(candidates) != 1:
        return None
    sheet = candidates[0]

    target = normalize_scheme_name(scheme_hint)
    if not target:
        return None
    for row in rows_for(sheet, limit=_TITLE_SCAN_ROWS):
        for c in row:
            if isinstance(c, str) and target in normalize_scheme_name(c):
                return sheet
    return None
