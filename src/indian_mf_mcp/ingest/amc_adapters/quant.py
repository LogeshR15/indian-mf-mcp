"""quant Mutual Fund adapter (quantmutual.com — lowercase "quant"; unrelated to Quantum AMC /
quantumamc.com).

**The prior BLOCKED verdict was stale, and it was recorded against the wrong page anyway.**
Whatever page was probed before ("empty ASP.NET WebForms viewstate, zero backing requests
fire") is not `https://quantmutual.com/statutory-disclosures` — the disclosure page named by
AMFI's own registry (https://www.amfiindia.com/online-center/portfolio-disclosure). Ground-
truthed live 2026-09-13: a plain `httpx.get()` with this project's honest User-Agent on that
exact URL returns a clean `200` with the full statutory-disclosures accordion already
rendered server-side — no viewstate problem, no dead page. The "MONTHLY PORTFOLIO" section
of that accordion, however, renders empty (`<div class='stat-cont-sub' id='MONTHLY
PORTFOLIO'></div>`) until a year tab is clicked; that is what looked broken. It isn't dead,
it's just AJAX-driven:

    POST https://quantmutual.com/statutorydisclosures.aspx/displaydisclouser
    Content-Type: application/json; charset=utf-8
    body: {"id": "<year>", "cat": "MONTHLY PORTFOLIO"}

This is a classic ASP.NET WebForms **PageMethod** (`[WebMethod]` static method on the
`.aspx` page, invoked via the standard `PageName.aspx/MethodName` convention) — not a SPA,
not a separate API host, and it needs no viewstate/session token at all: a bare, cookie-less
`httpx.post()` with a JSON body succeeds. The response is `{"d": "<ul>...</ul>"}` — a
JSON-string-escaped HTML fragment of `<li><a href='...'>Month Year</a></li>` entries, one per
month, for the clicked year. `id` accepts every year from 2018 (verified; the tab list itself
runs 2018-2026) through the current year — this is how history is reached: loop the adapter's
own `since` date across each year in range and issue one POST per year (9 requests worst
case for the full 2018-2026 span), rather than guessing a URL convention — there is none;
filenames are hand-uploaded and wildly inconsistent ("Monthly_Portfolio_Dec23.xlsx",
"DEC-21.xlsx", "Monthly Portfolio_August_2023.xlsx", "quant_Mutual_Fund_Monthly_Portfolio_
Jan2026.xlsx" all appear across different months). The one reliable field is each entry's own
anchor **text**, always "<Month name> <year>" (e.g. "December 2023", "August 2026"), which is
what this adapter parses for `as_of_date` — never the filename.

Older months (verified: 2021 August, and everything 2018-2020) are served as legacy `.xls`
(BIFF), not `.xlsx` — those are correctly discovered but rejected downstream by
`portfolio_ingest.py`'s format sniff (`skipped_format`), same as any other AMC's legacy
files; this adapter does not special-case them.

Each `.xlsx` is a **combined workbook**: one sheet per scheme, covering all ~28 quant
schemes for that month in a single file (SBI/Tata-style, not PPFAS/Union-style). There is,
however, **no separate "Index" lookup sheet** at all — sheet *names* are terse internal
acronyms ("qFLEXI", "qMES", "qL&MF", ...) that bear no resemblance to the scheme name, so
neither `combined_workbook.find_sheet_code` (needs an Index sheet) nor
`find_sheet_by_title` (needs the scheme name in each sheet's own row 1) applies as-is: row 1
of every sheet is the constant literal "quant Mutual Fund", identical across all schemes, and
the actual scheme name ("quant Flexi Cap Fund", "quant Small Cap Fund", ...) lives in row 2
instead. `_resolve_sheet` below is a small local variant of `find_sheet_by_title` reading row
2 rather than row 1 — kept local to this file rather than proposed as a shared-parser change,
since no other AMC seen so far needs a configurable title row.

Values are already fractional (e.g. 0.0963 = 9.63%), and the grand-total row is a plain
"GRAND TOTAL" — no other reconciliation quirks; parses against `xlsx_portfolio.py` with zero
changes (ground-truthed on quant Flexi Cap Fund, August 2026: 100% reconciliation, 29 rows
including section-header pseudo-rows, real holdings ISIN-complete).
"""
from __future__ import annotations

import html
import io
import json
import re
from datetime import date, datetime

import httpx
import openpyxl

from indian_mf_mcp import config
from indian_mf_mcp.ingest.amc_adapters.base import DocType, DocumentRef

DISCLOSURE_PAGE = "https://quantmutual.com/statutory-disclosures"
DISCLOSER_ENDPOINT = "https://quantmutual.com/statutorydisclosures.aspx/displaydisclouser"
BASE_URL = "https://quantmutual.com"
CATEGORY = "MONTHLY PORTFOLIO"

_ENTRY_RE = re.compile(r"<a href='([^']+)'[^>]*>([^<]+)</a>", re.IGNORECASE)
_TITLE_DATE_RE = re.compile(r"([A-Za-z]+)\s+(\d{4})")


def _years_since(since: date) -> list[int]:
    return list(range(since.year, date.today().year + 1))


def _resolve_sheet(raw: bytes, scheme_hint: str) -> str | None:
    """Local variant of combined_workbook.find_sheet_by_title: quant's sheets have no Index
    lookup sheet, and the scheme name lives in row 2 (row 1 is the constant "quant Mutual
    Fund" literal on every sheet, so row 1 alone cannot discriminate schemes)."""
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    target = scheme_hint.strip().lower()
    best = None
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            row2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        except StopIteration:
            continue
        title = next((c for c in row2 if isinstance(c, str) and c.strip()), None)
        if title is None:
            continue
        title_l = title.strip().lower()
        if title_l == target:
            return sheet_name
        if (target in title_l or title_l in target) and best is None:
            best = sheet_name
    return best


class QuantAdapter:
    amc_id = "amc-quant"

    def list_documents(self, doc_type: DocType, since: date,
                        scheme_hint: str | None = None, client: httpx.Client | None = None) -> list[DocumentRef]:
        if doc_type != DocType.MONTHLY_PORTFOLIO:
            return []
        headers = {"User-Agent": config.USER_AGENT, "Content-Type": "application/json; charset=utf-8"}
        post = client.post if client is not None else httpx.post

        refs: list[DocumentRef] = []
        seen = set()
        for year in _years_since(since):
            body = json.dumps({"id": str(year), "cat": CATEGORY})
            resp = post(DISCLOSER_ENDPOINT, headers=headers, content=body, timeout=45)
            resp.raise_for_status()
            try:
                fragment = resp.json().get("d", "")
            except ValueError:
                continue
            fragment = html.unescape(fragment)

            for path, title in _ENTRY_RE.findall(fragment):
                m = _TITLE_DATE_RE.search(title.strip())
                if not m:
                    continue
                month_name, year_str = m.groups()
                try:
                    as_of_month_start = datetime.strptime(f"{month_name} {year_str}", "%B %Y").date()
                except ValueError:
                    continue
                import calendar
                last_day = calendar.monthrange(as_of_month_start.year, as_of_month_start.month)[1]
                as_of = as_of_month_start.replace(day=last_day)
                if as_of < since:
                    continue
                url = path if path.startswith("http") else f"{BASE_URL}{path}"
                key = (url.split("?")[0], as_of)
                if key in seen:
                    continue
                seen.add(key)
                refs.append(DocumentRef(url=url, doc_type=DocType.MONTHLY_PORTFOLIO,
                                         as_of_date=as_of, scheme_hint=scheme_hint))
        refs.sort(key=lambda r: r.as_of_date)
        return refs

    def fetch(self, ref: DocumentRef, client: httpx.Client | None = None) -> bytes:
        headers = {"User-Agent": config.USER_AGENT}
        get = client.get if client is not None else httpx.get
        resp = get(ref.url, headers=headers, follow_redirects=True, timeout=60)
        resp.raise_for_status()
        return resp.content
