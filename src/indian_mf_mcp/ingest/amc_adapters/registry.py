"""Registry of AMC adapters actually wired up so far. Each entry can supply a `sheet_resolver`
for AMCs that publish one combined workbook per month (select the scheme's own sheet) —
None means one file per scheme (the default/PPFAS-style layout).

Adding an AMC here means: (1) a real, live-verified way to discover its monthly portfolio
files (a static link, or a documented backing endpoint found via a one-time build-time
discovery, e.g. Playwright network capture — never a headless browser at runtime), and
(2) confirmation the xlsx_portfolio parser's column/header/percentage-scale detection
handles that AMC's layout (add a golden fixture test before trusting reconciliation here).
"""
from __future__ import annotations

import io
from typing import Callable

import openpyxl

from indian_mf_mcp.ingest.amc_adapters import combined_workbook
from indian_mf_mcp.ingest.amc_adapters.axis import AxisAdapter
from indian_mf_mcp.ingest.amc_adapters.bank_of_india import BankOfIndiaAdapter
from indian_mf_mcp.ingest.amc_adapters.baroda_bnp_paribas import BarodaBNPParibasAdapter
from indian_mf_mcp.ingest.amc_adapters.dsp import DSPAdapter
from indian_mf_mcp.ingest.amc_adapters.franklin_templeton import FranklinTempletonAdapter
from indian_mf_mcp.ingest.amc_adapters.hdfc import HDFCAdapter
from indian_mf_mcp.ingest.amc_adapters.lic import LicAdapter
from indian_mf_mcp.ingest.amc_adapters.mirae import MiraeAdapter
from indian_mf_mcp.ingest.amc_adapters.motilal_oswal import MotilalOswalAdapter
from indian_mf_mcp.ingest.amc_adapters.navi import NaviAdapter
from indian_mf_mcp.ingest.amc_adapters.nippon import NipponAdapter
from indian_mf_mcp.ingest.amc_adapters.ppfas import PPFASAdapter
from indian_mf_mcp.ingest.amc_adapters.quant import QuantAdapter
from indian_mf_mcp.ingest.amc_adapters.quant import _resolve_sheet as _quant_sheet_resolver
from indian_mf_mcp.ingest.amc_adapters.sbi import SBIAdapter
from indian_mf_mcp.ingest.amc_adapters.sundaram import SundaramAdapter
from indian_mf_mcp.ingest.amc_adapters.tata import TataAdapter
from indian_mf_mcp.ingest.amc_adapters.taurus import TaurusAdapter
from indian_mf_mcp.ingest.amc_adapters.union import UnionAdapter
from indian_mf_mcp.ingest.amc_adapters.uti import UTIAdapter
from indian_mf_mcp.ingest.amc_adapters.zerodha import ZerodhaAdapter


def _sbi_sheet_resolver(raw: bytes, scheme_hint: str) -> str | None:
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    if "Index" not in wb.sheetnames:
        return None
    index_rows = list(wb["Index"].iter_rows(values_only=True))
    return SBIAdapter.find_sheet_code(index_rows, scheme_hint)


ADAPTERS: dict[str, tuple[type, Callable[[bytes, str], str | None] | None]] = {
    "ppfas": (PPFASAdapter, None),
    "sbi": (SBIAdapter, _sbi_sheet_resolver),
    "uti": (UTIAdapter, None),
    "mirae": (MiraeAdapter, None),
    "motilal-oswal": (MotilalOswalAdapter, combined_workbook.find_sheet_code),
    "tata": (TataAdapter, combined_workbook.find_sheet_code),
    "nippon": (NipponAdapter, combined_workbook.find_sheet_code),
    "dsp": (DSPAdapter, None),
    "franklin-templeton": (FranklinTempletonAdapter, combined_workbook.find_sheet_by_title),
    "baroda-bnp-paribas": (BarodaBNPParibasAdapter, combined_workbook.find_sheet_code),
    "sundaram": (SundaramAdapter, combined_workbook.find_sheet_code),
    "union": (UnionAdapter, None),
    "lic": (LicAdapter, None),
    "taurus": (TaurusAdapter, None),
    "bank-of-india": (BankOfIndiaAdapter, combined_workbook.find_sheet_code),
    "hdfc": (HDFCAdapter, None),
    "quant": (QuantAdapter, _quant_sheet_resolver),
    "navi": (NaviAdapter, None),
    "zerodha": (ZerodhaAdapter, None),
    "axis": (AxisAdapter, None),
}


def get_adapter(amc_key: str):
    """Returns (adapter_instance, sheet_resolver_or_None)."""
    adapter_cls, resolver = ADAPTERS[amc_key]
    return adapter_cls(), resolver
