"""Parser for AMC monthly/fortnightly portfolio disclosure XLSX files.

Ground-truthed against a real file: PPFAS Flexi Cap, August 2026
(PPFCF_PPFAS_Monthly_Portfolio_Report_August_31_2026.xlsx), fetched live 2026-09-12.

Design rules (spec §7):
  - Do not assume a fixed header row index — scan for the row containing "ISIN" + "% to Net".
  - Track a running section label to assign asset_class; do not flatten sub-sections
    (e.g. Reits, Arbitrage, Foreign Investments, Derivatives all carry distinct conventions).
  - Reconciliation: the GRAND TOTAL row's %-to-NAV must be ~1.0 (fractional). If it deviates
    beyond tolerance, flag the snapshot rather than silently serving a broken portfolio.
  - Derivatives use a *different* column layout in the same file and must never be netted
    into equity exposure.
"""
from __future__ import annotations

import datetime as _dt
import io
import re
from dataclasses import dataclass, field

import openpyxl
import xlrd

# Keyword substrings (not exact phrases — wording varies by AMC, e.g. PPFAS "Equity & Equity
# Related" vs UTI "EQUITY AND EQUITY RELATED") that mark a header row as TOP-LEVEL, resetting
# the top_section tracker. Any other bare-text header row (e.g. "(a) Listed...", "Arbitrage",
# "Certificate of Deposit") is treated as a sub-section label within the active top_section.
TOP_LEVEL_SECTION_KEYWORDS = (
    "equity & equity related", "equity and equity related", "equity related instruments",
    "money market instruments", "debt instruments", "government securities", "derivatives",
    "others",
)

AGGREGATE_ROW_LABELS = {"sub total", "total", "grand total"}

# Footer markers where structured holdings/derivatives data ends and free-text disclosures,
# performance tables, and quant indicators begin. Parsing stops here to avoid misreading
# footer numbers (returns tables, ratios) as spurious holding/derivative rows.
FOOTER_STOP_MARKERS = (
    "# traded", "notes & symbols", "notes:", "lumpsum investment performance",
    "sip investment performance", "quantitative indicators", "this product is suitable",
    "total below investment grade", "details of intra scheme investments",
    "hedging positions", "nav as on",
    # HSBC's debt-scheme files print these two quant-indicator lines loose, with no
    # "Quantitative Indicators" section header above them to trip the existing marker.
    "annualised portfolio ytm", "macaulay duration",
)

# Some AMCs (Tata) label the true fund-level 100% total "NET ASSETS" rather than "GRAND
# TOTAL" or anything containing the word "total" at all — must be recognised explicitly or
# it falls through as a spurious duplicate holding worth ~100% of NAV. HSBC suffixes its own
# version with the as-of date ("Total Net Assets as on 31-August-2026"), so this is matched as
# a prefix, not an exact label — a plain "total"-containing row would otherwise be discarded
# by the aggregate-row skip below with no total ever recorded.
_EXPLICIT_GRAND_TOTAL_LABELS = {"net assets", "total net assets"}

# HSBC's own title cell says "as of" rather than every other AMC's "as on" — a wording
# variant, so it is accepted alongside "as on" rather than handled per-AMC.
AS_OF_RE = re.compile(r"as o[nf]\s+([A-Za-z]+ \d{1,2},?\s*\d{4})", re.IGNORECASE)
AS_OF_NUMERIC_RE = re.compile(r"as o[nf]\s+(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})", re.IGNORECASE)
BENCHMARK_RE = re.compile(r"\(([A-Za-z0-9&/ ]*\bTRI\b[A-Za-z0-9&/ ]*)\)")


@dataclass
class ParsedHolding:
    instrument_name: str
    isin: str | None
    industry_or_rating: str | None
    quantity: float | None
    market_value_lakhs: float | None
    pct_nav: float | None  # fractional, e.g. 0.0763 = 7.63%
    asset_class: str
    listed: bool
    section_label: str


@dataclass
class ParsedDerivative:
    instrument_name: str
    direction: str | None       # "Long" | "Short" | None
    quantity: float | None
    market_value_lakhs: float | None
    pct_to_aum: float | None
    section_label: str


@dataclass
class PortfolioParseResult:
    holdings: list[ParsedHolding] = field(default_factory=list)
    derivatives: list[ParsedDerivative] = field(default_factory=list)
    grand_total_market_value: float | None = None
    grand_total_pct_nav: float | None = None
    reconciliation_ok: bool | None = None  # None if GRAND TOTAL not found at all
    as_of_date_str: str | None = None      # raw string as printed, e.g. "August 31, 2026"
    benchmark_name: str | None = None
    sheet_count: int = 0  # total sheets in the source workbook, regardless of which was parsed
    parse_confidence: float = 0.0
    warnings: list[str] = field(default_factory=list)


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _is_blank(value) -> bool:
    """Empty for header-detection purposes. Some AMCs (ICICI Prudential) leave unused ISIN /
    industry / quantity cells as an empty string rather than None."""
    return value is None or (isinstance(value, str) and not value.strip())


_TOTAL_WORD_RE = re.compile(r"\b(?:sub)?total\b", re.IGNORECASE)


def _is_aggregate_label(text: str) -> bool:
    """Match "total" as a whole word ANYWHERE in the label, not just as a prefix — AMCs are
    inconsistent about where they put it: UTI prefixes ("TOTAL:(a) Listed...", "TOTAL :
    <Scheme Name>"), Tata suffixes ("EQUITY & EQUITY RELATED TOTAL", "PORTFOLIO TOTAL"). A
    prefix-only or exact-match check silently lets suffix-style totals through as fake
    holdings (they have no ISIN/industry/quantity, only a value+pct, identical in shape to a
    genuine standalone cash line like "NET CURRENT ASSETS" — the label text is the only
    signal that distinguishes them). Nippon India writes "Subtotal" as one word, which a bare
    \\btotal\\b never matches — it was loaded as a ~99%-of-NAV fake holding.

    Callers must only apply this to rows WITHOUT an ISIN: "total" is also a word inside real
    company names ("Adani Total Gas Ltd."), and skipping that row drops a genuine holding."""
    return bool(_TOTAL_WORD_RE.search(text))


# Leading "(a) ", "a) ", "A. ", "II. " enumerators on section labels.
_ENUMERATOR_RE = re.compile(r"^\(?[a-z0-9]{1,4}[).:]\s*")


def _nil_to_none(value):
    """Several AMCs (ICICI Prudential, Mirae) print the literal text "Nil" in the value/%
    columns of an empty section. That's "no value", not a value — left as a string it
    defeats the "no data columns populated" section-header test and the empty section is
    stored as a fake zero-weight holding."""
    if isinstance(value, str) and value.strip().lower() in ("nil", ""):
        return None
    return value


def _asset_class_for(top_section: str, sub_section: str) -> str:
    t = top_section.lower()
    s = sub_section.lower()
    if "foreign" in t:
        return "foreign"
    if "reits" in s:
        return "reit"
    if "equity" in t:
        return "equity"
    if "money market" in t or "debt instruments" in t or "government securities" in t:
        return "debt"
    if "reverse repo" in s or "treps" in s:
        return "cash"
    if "mutual fund" in s:
        return "cash"
    if "derivatives" in t:
        return "derivative"
    return "other"


@dataclass
class ColumnMap:
    name: int
    isin: int
    industry: int
    quantity: int
    value: int
    pct: int
    # Set only when the file has a SEPARATE rating column alongside industry (quant: "RATING"
    # then "INDUSTRY"); a combined "Industry/Rating" column is just `industry`.
    rating: int | None = None


def _find_main_header_row(rows: list[tuple]) -> tuple[int, ColumnMap] | None:
    """Header wording AND column position both differ by AMC (PPFAS: name at col 1, "%
    to Net Assets"; SBI: name at col 2, "% to AUM", one column further right because SBI
    also carries a scheme/instrument code column). Detect each column's position from the
    header row's own text rather than assuming a fixed offset (spec §7)."""
    for i, row in enumerate(rows):
        texts = [str(c).strip().lower() if isinstance(c, str) else "" for c in row]

        def find(*needles: str) -> int | None:
            for idx, t in enumerate(texts):
                if any(t.startswith(n) or n in t for n in needles):
                    return idx
            return None

        name_col = find("name of the instrument", "name of instrument", "instrument name")
        isin_col = find("isin")
        # HSBC spells this out in full ("Percentage to Net Assets") instead of using the "%"
        # symbol every other AMC observed so far uses — a wording variant, not a layout quirk,
        # so it belongs here alongside the other synonyms rather than in HSBC's own adapter.
        pct_col = find("% to net", "% to aum", "% to nav", "% of net",
                        "percentage to net", "percentage to aum", "percentage to nav",
                        "percentage of net")
        if name_col is None or isin_col is None or pct_col is None:
            continue
        # Prefer the industry column when a file carries both: quant lists "RATING" (always
        # "N.A." for equities) BEFORE "INDUSTRY", and a first-match on either word read every
        # equity's sector as "N.A.". The rating column is kept as a per-row fallback so debt
        # rows (industry "N.A.", rating "SOV") still get their rating.
        industry_col = find("industry")
        rating_col = find("rating")
        if industry_col is None:
            industry_col, rating_col = rating_col, None
        elif rating_col == industry_col:
            rating_col = None
        qty_col = find("quantity")
        value_col = next(
            (idx for idx, t in enumerate(texts)
             if ("market" in t and "value" in t) or ("mkt" in t and "val" in t)),
            None,
        )
        if industry_col is None or qty_col is None or value_col is None:
            continue
        return i, ColumnMap(name=name_col, isin=isin_col, industry=industry_col,
                             quantity=qty_col, value=value_col, pct=pct_col, rating=rating_col)
    return None


def _find_as_of(rows: list[tuple]) -> str | None:
    import datetime as _dt
    for row in rows[:10]:
        for i, c in enumerate(row):
            if isinstance(c, str):
                m = AS_OF_RE.search(c)
                if m:
                    return m.group(1).strip()
                m2 = AS_OF_NUMERIC_RE.search(c)
                if m2:
                    return m2.group(1).strip()
                if "as on" in c.lower():
                    # some AMCs (e.g. SBI) put a real datetime in an adjacent cell rather
                    # than embedding the date in the label string itself.
                    for other in row[i + 1:]:
                        if isinstance(other, _dt.datetime):
                            return f"{other.strftime('%B')} {other.day}, {other.year}"
    return None


def _find_benchmark(rows: list[tuple]) -> str | None:
    """Prefer the riskometer-adjacent benchmark name (authoritative footer location per spec
    §3.1) over any earlier stray "(...TRI)" occurrence in a returns table."""
    for i, row in enumerate(rows):
        for cell in row:
            if isinstance(cell, str) and "benchmark" in cell.lower() and "riskometer" in cell.lower():
                for follow_row in rows[i : i + 3]:
                    for fcell in follow_row:
                        if isinstance(fcell, str):
                            m = BENCHMARK_RE.search(fcell)
                            if m:
                                return m.group(1).strip()
    for row in rows:
        for cell in row:
            if isinstance(cell, str):
                m = BENCHMARK_RE.search(cell)
                if m:
                    return m.group(1).strip()
    return None


def parse_portfolio_xlsx(raw: bytes, sheet_name: str | None = None) -> PortfolioParseResult:
    """sheet_name: for AMCs that publish one combined workbook covering every scheme (e.g.
    SBI: one sheet per scheme, an "Index" sheet mapping short-codes to names), select the
    scheme's own sheet. None uses the first sheet (PPFAS-style: one file per scheme)."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced via parse_confidence, not raised
        return _open_failure(f"workbook failed to open: {exc}")

    if sheet_name is not None:
        if sheet_name not in wb.sheetnames:
            return _open_failure(f"sheet {sheet_name!r} not found in workbook "
                                  f"(available: {wb.sheetnames[:10]}...)")
        ws = wb[sheet_name]
    else:
        ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    result = _parse_rows(rows)
    result.sheet_count = len(wb.sheetnames)
    return result


def parse_portfolio_xls(raw: bytes, sheet_name: str | None = None) -> PortfolioParseResult:
    """Legacy binary Excel (.xls / BIFF) counterpart of parse_portfolio_xlsx — same AMC
    layout conventions and column-detection heuristics, different container format (spec §7:
    "XLS (legacy BIFF) | xlrd<2.0 or LibreOffice headless conversion"). In practice xlrd>=2.0
    dropped .xlsx support entirely and reads only .xls now, which is exactly the format this
    function targets, so the project's xlrd>=2.0.1 pin is already the right one — no upper
    bound needed. Shares the entire row-walking core with parse_portfolio_xlsx via
    _parse_rows(); only raw-bytes-to-rows extraction differs by format."""
    try:
        wb = xlrd.open_workbook(file_contents=raw)
    except Exception as exc:  # noqa: BLE001 - surfaced via parse_confidence, not raised
        return _open_failure(f"workbook failed to open: {exc}")

    sheet_names = wb.sheet_names()
    if sheet_name is not None:
        if sheet_name not in sheet_names:
            return _open_failure(f"sheet {sheet_name!r} not found in workbook "
                                  f"(available: {sheet_names[:10]}...)")
        ws = wb.sheet_by_name(sheet_name)
    else:
        ws = wb.sheet_by_index(0)
    rows = xls_sheet_to_rows(ws)
    result = _parse_rows(rows)
    result.sheet_count = len(sheet_names)
    return result


def _open_failure(message: str) -> PortfolioParseResult:
    result = PortfolioParseResult()
    result.warnings.append(message)
    result.parse_confidence = 0.0
    return result


def xls_sheet_to_rows(ws) -> list[tuple]:
    """Convert an xlrd Sheet into the same list[tuple] shape openpyxl's
    `ws.iter_rows(values_only=True)` produces, so _parse_rows() never needs to know which
    library produced its input. Two format differences are normalised here rather than in
    the shared row-walking logic: xlrd represents a blank cell as '' (empty string), not
    None — openpyxl's convention, which _is_blank()/_num() already assume — and a date cell
    as a raw float serial rather than a datetime object."""
    rows: list[tuple] = []
    for r in range(ws.nrows):
        row = []
        for c in range(ws.ncols):
            cell = ws.cell(r, c)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = _dt.datetime(*xlrd.xldate_as_tuple(value, ws.book.datemode))
                except Exception:  # noqa: BLE001 - fall back to the raw serial, never raise
                    pass
            elif value == "":
                value = None
            row.append(value)
        rows.append(tuple(row))
    return rows


def _drop_subtotal_headers(result: PortfolioParseResult, standalone_idx: list[int]) -> None:
    """ICICI Prudential prints a section's subtotal on the section header row itself
    ("Treasury Bills ... 12697.96  0.0049", then the individual T-bills). That row has the same
    shape as a genuine standalone line (TREPS, NET CURRENT ASSETS — value + %, no ISIN), so it
    can't be told apart when it is read. It can afterwards: a header is followed by rows (at
    least one with an ISIN) that add up exactly to its own %, a standalone line is not.
    Anything that doesn't prove itself a header this way stays a holding."""
    drop = set()
    bounds = standalone_idx[1:] + [len(result.holdings)]
    for i, end in zip(standalone_idx, bounds):
        header = result.holdings[i]
        if header.pct_nav is None:
            continue
        # Children run to the next standalone line, not the next ISIN-less row: ICICI puts
        # hedge legs ("Maruti Suzuki India Ltd. $$", no ISIN, negative %) inside the block,
        # and the header's subtotal includes them.
        children = result.holdings[i + 1:end]
        if not any(h.isin for h in children):
            continue
        child_sum = sum(h.pct_nav or 0.0 for h in children)
        if abs(child_sum - header.pct_nav) <= max(1e-6, 1e-4 * abs(header.pct_nav)):
            drop.add(i)
            for h in children:
                h.section_label = header.instrument_name
                name_l = header.instrument_name.lower()
                h.listed = "listed" in name_l and "unlisted" not in name_l
    if drop:
        result.holdings = [h for i, h in enumerate(result.holdings) if i not in drop]


def _parse_rows(rows: list[tuple]) -> PortfolioParseResult:
    """Shared row-walking core for parse_portfolio_xlsx and parse_portfolio_xls — everything
    from here down operates on a plain list[tuple] and has no idea which spreadsheet library
    or file format produced it."""
    result = PortfolioParseResult()
    result.as_of_date_str = _find_as_of(rows)
    result.benchmark_name = _find_benchmark(rows)

    found = _find_main_header_row(rows)
    if found is None:
        result.warnings.append("could not locate main holdings header row (ISIN + % to Net/AUM)")
        result.parse_confidence = 0.0
        return result
    header_idx, cols = found
    # Derivatives table columns are the same relative offsets after the name column, one
    # position earlier than the main table (direction where the main table has ISIN, etc.) —
    # true for every AMC observed so far (PPFAS, SBI); re-verified per-file via deriv_header_seen.
    deriv_direction_col = cols.name + 1
    deriv_qty_col = cols.name + 2
    deriv_value_col = cols.name + 3
    deriv_pct_col = cols.name + 4

    top_section = ""
    sub_section = ""
    in_derivatives_body = False
    deriv_header_seen = False
    # Indices into result.holdings of valued rows with no ISIN/industry/quantity — either a
    # genuine standalone line (NET CURRENT ASSETS, TREPS) or a section header that carries its
    # own subtotal (see _drop_subtotal_headers). Resolved after the walk.
    standalone_idx: list[int] = []
    # The GRAND TOTAL row closes the holdings table. What follows can still be a derivatives
    # exposure table (PPFAS), but never more holdings: Franklin prints an interest-rate-swap
    # notional table and the NAV-per-plan block there, which were being loaded as ~16% of
    # phantom NAV.
    past_grand_total = False

    def cell(row, idx):
        return row[idx] if idx is not None and len(row) > idx else None

    for row in rows[header_idx + 1:]:
        col_b = cell(row, cols.name)
        col_c_raw = cell(row, cols.isin)
        col_d = cell(row, cols.industry)
        col_e = cell(row, cols.quantity)
        col_f = _nil_to_none(cell(row, cols.value))
        col_g = _nil_to_none(cell(row, cols.pct))

        # Some AMCs (Franklin Templeton) put ISIN *before* the name column, and for
        # section/total/summary rows the label text sits in that leading ISIN-column
        # position instead of the name column (since there's no ISIN to put there) —
        # fall back to whichever of the two actually holds text. When col_c is consumed
        # as the label fallback, treat it as empty for ISIN/data-column purposes below —
        # it was never a real ISIN, so it must not defeat the "no data columns populated"
        # section-header check or get stored as a bogus ISIN on a fake holding.
        label_from_name = str(col_b).strip() if isinstance(col_b, str) else None
        used_isin_as_label = label_from_name is None and isinstance(col_c_raw, str)
        label = label_from_name or (str(col_c_raw).strip() if used_isin_as_label else None)
        col_c = None if used_isin_as_label else col_c_raw

        if label and any(label.lower().startswith(m) for m in FOOTER_STOP_MARKERS):
            break

        # SEBI's mandated scheme-riskometer block (added across many AMCs' files) repeats the
        # scheme's own name as a row label with "Scheme Riskometer" / "Scheme Benchmark
        # Riskometer" in the industry/value columns instead of real holding data — the
        # scheme-name label defeats FOOTER_STOP_MARKERS (it never starts with a known phrase),
        # so it must be caught on the industry column's own fixed text instead (HSBC: this
        # follows directly after the GRAND TOTAL row with nothing valid after it).
        if isinstance(col_d, str) and "riskometer" in col_d.lower():
            break

        if label and (label.strip().lower().startswith("grand total")
                      or any(label.strip().lower().startswith(t)
                             for t in _EXPLICIT_GRAND_TOTAL_LABELS)):
            result.grand_total_market_value = _num(col_f)
            result.grand_total_pct_nav = _num(col_g)
            past_grand_total = True
            continue

        if label and label.strip().lower() == "nil":
            continue  # placeholder row for an empty section (e.g. "(B) Unlisted ... Nil")

        # Any "*total*" row (prefix or suffix style — see _is_aggregate_label) is a
        # sub-total/aggregate, not a holding: skip it. Some AMCs (UTI, Tata) never print an
        # explicit "GRAND TOTAL" %-figure at all; the fallback below self-sums the real
        # holdings instead of trying to guess which aggregate row is the true fund-level one.
        if label and _is_blank(col_c) and _is_aggregate_label(label):
            continue

        if label and label.strip().lower() == "derivatives":
            in_derivatives_body = True
            top_section = label
            sub_section = ""
            continue

        no_identity_data = _is_blank(col_c) and _is_blank(col_d) and _is_blank(col_e)
        looks_like_top_level = bool(label) and any(
            kw in label.strip().lower() for kw in TOP_LEVEL_SECTION_KEYWORDS
        )
        # A row carrying its own value is only a top-level header if it actually STARTS with
        # the section name (ICICI's "Money Market Instruments  12697.96  0.0049"). Anywhere-
        # in-label matching turned Nippon's valued "Cash Margin - Derivatives" line (0.5% of
        # NAV) into a "derivatives" section header and dropped it.
        if looks_like_top_level and (_num(col_f) is not None or _num(col_g) is not None):
            bare = _ENUMERATOR_RE.sub("", label.strip().lower())
            looks_like_top_level = any(bare.startswith(kw) for kw in TOP_LEVEL_SECTION_KEYWORDS)

        # A DERIVATIVES block is not always the last section of the holdings table: quant,
        # Baroda BNP Paribas and Bajaj Finserv follow it with DEBT / MONEY MARKET / OTHERS
        # (T-bills, TREPS, net current assets). Staying in derivatives mode swallowed those
        # rows as derivatives, dropping up to ~29% of NAV from holdings.
        if in_derivatives_body and looks_like_top_level and no_identity_data:
            in_derivatives_body = False

        if in_derivatives_body:
            # Two derivative layouts exist. PPFAS/SBI print a separate sub-table with its own
            # "Long/Short" header, one column left of the main table. quant prints futures
            # as ordinary rows of the main table (contract symbol in the ISIN column, % to NAV
            # in the main % column; Baroda BNP Paribas the same with a blank ISIN) with no
            # sub-header — read with the offset layout, quantity landed in market value and
            # market value in % (334% of AUM).
            main_layout = (not deriv_header_seen and _num(col_f) is not None
                           and _num(col_g) is not None)
            if main_layout:
                deriv_direction, deriv_qty, deriv_value, deriv_pct = None, col_e, col_f, col_g
            else:
                deriv_direction = cell(row, deriv_direction_col)
                deriv_qty = cell(row, deriv_qty_col)
                deriv_value = cell(row, deriv_value_col)
                deriv_pct = cell(row, deriv_pct_col)
            if label and not deriv_header_seen and deriv_direction and \
                    "long" in str(deriv_direction).strip().lower():
                deriv_header_seen = True
                continue
            if label and not main_layout and _is_aggregate_label(label):
                continue
            if label and not main_layout and deriv_direction is None and deriv_value is None:
                # sub-section header within derivatives, e.g. "Index / Stock Futures"
                sub_section = label
                continue
            if label:
                result.derivatives.append(ParsedDerivative(
                    instrument_name=label,
                    direction=str(deriv_direction).strip("() ") if isinstance(deriv_direction, str) else None,
                    quantity=_num(deriv_qty), market_value_lakhs=_num(deriv_value),
                    pct_to_aum=_num(deriv_pct),
                    section_label=sub_section or top_section,
                ))
            continue

        # Non-derivatives body. A row is a section/sub-section header only if NONE of the
        # data columns are populated — some AMCs (e.g. UTI's "NET CURRENT ASSETS") report a
        # standalone cash/other line with a real value+pct but no ISIN/industry/quantity;
        # treating that as a header would silently drop it from holdings and reconciliation.
        # A named top-level section stays a section even when it also carries its own
        # aggregate value/pct on the same row (ICICI Prudential does this). Anything else
        # still needs every data column empty, which is what keeps a standalone valued line
        # like UTI's "NET CURRENT ASSETS" classified as a holding rather than a header.
        if label and no_identity_data and (
            looks_like_top_level or (col_f is None and col_g is None)
        ):
            if looks_like_top_level:
                top_section = label
                sub_section = ""
            else:
                sub_section = label
            continue

        if label and past_grand_total:
            continue

        if label:
            isin = col_c.strip() if isinstance(col_c, str) and col_c.strip() else None
            industry = col_d.strip() if isinstance(col_d, str) else None
            if cols.rating is not None and (not industry or industry.lower() in ("n.a.", "na", "-")):
                rating = cell(row, cols.rating)
                if isinstance(rating, str) and rating.strip():
                    industry = rating.strip()
            if no_identity_data:
                standalone_idx.append(len(result.holdings))
            label_l = label.lower()
            # Cash-equivalent line items are labeled this way across AMCs regardless of
            # whatever top/sub-section they happen to sit under in that file.
            if any(k in label_l for k in ("net current assets", "net receivables", "cash & cash equivalent")):
                asset_class = "cash"
            else:
                asset_class = _asset_class_for(top_section, sub_section)
            result.holdings.append(ParsedHolding(
                instrument_name=label,
                isin=isin,
                industry_or_rating=industry,
                quantity=_num(col_e),
                market_value_lakhs=_num(col_f),
                pct_nav=_num(col_g),
                asset_class=asset_class,
                listed="listed" in sub_section.lower() if sub_section else False,
                section_label=sub_section or top_section,
            ))

    _drop_subtotal_headers(result, standalone_idx)
    file_stated_total = result.grand_total_pct_nav is not None

    # Some AMCs (e.g. UTI) never print a %-to-NAV figure on their grand-total row at all —
    # only a market value. Fall back to self-summing the extracted holdings' own pct_nav as
    # the reconciliation target, clearly labeled as self-derived rather than the file's own
    # stated total (a weaker but still useful integrity check).
    if result.grand_total_pct_nav is None and result.holdings:
        self_summed = sum(h.pct_nav for h in result.holdings if h.pct_nav is not None)
        if self_summed > 0:
            result.grand_total_pct_nav = self_summed
            result.warnings.append(
                f"file did not state a %-to-NAV grand total; using self-summed holdings "
                f"pct_nav ({self_summed}) as the reconciliation target instead.")

    # %-to-NAV convention differs by AMC: PPFAS reports fractional (0.0763 = 7.63%), SBI
    # reports percentage points (7.22 = 7.22%). Detect from the GRAND TOTAL scale and
    # normalise everything to fractional so downstream analytics never see mixed units.
    if result.grand_total_pct_nav is not None and result.grand_total_pct_nav > 5:
        scale = 100.0
        result.warnings.append(
            f"pct_nav values appear to be in percentage points (grand total="
            f"{result.grand_total_pct_nav}); normalised to fractional by dividing by 100.")
        result.grand_total_pct_nav = result.grand_total_pct_nav / scale
        for h in result.holdings:
            if h.pct_nav is not None:
                h.pct_nav = h.pct_nav / scale
        for d in result.derivatives:
            if d.pct_to_aum is not None:
                d.pct_to_aum = d.pct_to_aum / scale

    if result.grand_total_pct_nav is not None:
        # spec §7: sum of "% to Net Assets" must land within ~0.5% of 100 or the snapshot
        # is flagged rather than silently served as reconciled.
        result.reconciliation_ok = abs(result.grand_total_pct_nav - 1.0) <= 0.005
        # The file's own GRAND TOTAL saying 100% proves nothing about what WE extracted. Also
        # require the extracted rows to add up to it — otherwise a dropped section (quant:
        # 71%) or a double-counted subtotal (Nippon: 199%) sails through as "reconciled".
        # Futures are part of NAV in some layouts (quant) and a separate exposure table in
        # others (PPFAS), so either holdings alone or holdings + derivatives may match.
        if result.reconciliation_ok and file_stated_total:
            held = sum(h.pct_nav for h in result.holdings if h.pct_nav is not None)
            derived = held + sum(d.pct_to_aum for d in result.derivatives if d.pct_to_aum is not None)
            target = result.grand_total_pct_nav
            if abs(held - target) > 0.005 and abs(derived - target) > 0.005:
                result.reconciliation_ok = False
                result.warnings.append(
                    f"extracted rows sum to {held:.4f} of NAV ({derived:.4f} incl. derivatives) "
                    f"but the file's grand total is {target:.4f}; parse is incomplete or "
                    f"double-counted")
    else:
        result.reconciliation_ok = None
        result.warnings.append("GRAND TOTAL row not found; cannot verify 100% reconciliation")

    result.parse_confidence = (
        1.0 if result.reconciliation_ok else (0.5 if result.reconciliation_ok is False else 0.3)
    )
    if not result.holdings:
        result.parse_confidence = 0.0
        result.warnings.append("no holdings extracted")

    return result
